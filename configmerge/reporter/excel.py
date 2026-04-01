"""
configmerge.reporter.excel
~~~~~~~~~~~~~~~~~~~~~~~~~~
Generates the Excel (.xlsx) report with 6 sheets:

  1. MergeChanges       — all parameter-level changes
  2. BaseOnlyFiles      — files in base with no release counterpart
  3. ReleaseOnlyFiles   — files in release with no base counterpart
  4. FileMappings       — base ↔ release filename mappings applied
  5. ExcludedBaseOnly   — params skipped via --exclude-params-in-baseonlyconfig
  6. BaseConfigAsIs     — files copied as-is from base (--copy-baseonlyconfigfile)
"""

from __future__ import annotations

import os
import re
import json
import xml.etree.ElementTree as ET
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import List, Set, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill

from ..models import MergeConfig, MergeResult, ReportEntry, EntryType


# ---------------------------------------------------------------------------
# Colour palette
# ---------------------------------------------------------------------------

class _C:
    HEADER_FILL  = "1E2A3A"
    HEADER_FONT  = "FFFFFF"
    RED          = "C62828"
    ORANGE       = "E65100"
    YELLOW_BG    = "FFF9C4"
    GREEN        = "1B5E20"
    GREEN_BG     = "E8F5E9"
    ORANGE_BG    = "FFF3E0"
    RED_BG       = "FFEBEE"
    NORMAL       = "222222"


# ---------------------------------------------------------------------------
# Diff helpers
# ---------------------------------------------------------------------------

def _extract_xml_changes(old_xml: str, new_xml: str) -> str:
    try:
        old_root = ET.fromstring(f"<root>{old_xml}</root>")
        new_root = ET.fromstring(f"<root>{new_xml}</root>")
    except Exception:
        try:
            old_root = ET.fromstring(old_xml)
            new_root = ET.fromstring(new_xml)
        except Exception:
            return ""

    changes = []

    def build_map(root):
        m = defaultdict(list)
        for elem in root.iter():
            tag  = elem.tag.split("}")[-1]
            name = elem.attrib.get("name", "")
            m[(tag, name)].append(elem)
        return m

    old_map = build_map(old_root)
    new_map = build_map(new_root)

    for key in set(old_map) | set(new_map):
        tag, name = key
        old_elems = old_map.get(key, [])
        new_elems = new_map.get(key, [])

        if old_elems and not new_elems:
            changes.append(f"REMOVED: <{tag}> {name}")
            continue
        if new_elems and not old_elems:
            changes.append(f"ADDED: <{tag}> {name}")
            continue

        oe = old_elems[0]
        ne = new_elems[0]

        # Text change
        ot = (oe.text or "").strip()
        nt = (ne.text or "").strip()
        if ot != nt:
            changes.append(f"<{tag}> text: {ot!r} → {nt!r}")

        # Attribute change
        for attr in set(oe.attrib) | set(ne.attrib):
            ov = oe.attrib.get(attr, "")
            nv = ne.attrib.get(attr, "")
            if ov != nv:
                changes.append(f"<{tag}> @{attr}: {ov!r} → {nv!r}")

    return "\n".join(changes)


def _extract_json_changes(old_str: str, new_str: str) -> str:
    try:
        old = json.loads(old_str)
        new = json.loads(new_str)
    except Exception:
        return ""
    if old == new:
        return ""
    return f"OLD: {json.dumps(old)}\n→ NEW: {json.dumps(new)}"


def _build_changes_column(entry: ReportEntry) -> str:
    t = entry.type
    if t == EntryType.XML_BASE_TO_RELEASE_REPLACED:
        return _extract_xml_changes(entry.old, entry.new)
    if t == EntryType.JSON_BASE_TO_RELEASE_REPLACED:
        return _extract_json_changes(entry.old, entry.new)
    if t in (EntryType.BASE_TO_RELEASE_REPLACED, EntryType.UNCOMMENT_REPLACE):
        return f"release: {entry.old!r}  →  merged: {entry.new!r}"
    if t == EntryType.COMMA_VALUE_UNION:
        return f"release: {entry.old}\nmerged: {entry.new}"
    return ""


# ---------------------------------------------------------------------------
# Sheet builder
# ---------------------------------------------------------------------------

def _thin_border() -> Border:
    side = Side(style="thin")
    return Border(left=side, right=side, top=side, bottom=side)


def _add_sheet(
    wb: Workbook,
    name: str,
    headers: List[str],
    rows: List[List],
    row_colors: dict = None,  # {row_type_str: (font_color, fill_color)}
) -> None:
    ws = wb.create_sheet(name)
    border  = _thin_border()
    wrap    = Alignment(wrap_text=True, vertical="top")
    bold    = Font(bold=True, color=_C.HEADER_FONT)
    hdr_fill = PatternFill("solid", fgColor=_C.HEADER_FILL)

    # Header row
    ws.append(headers)
    for cell in ws[1]:
        cell.font      = bold
        cell.fill      = hdr_fill
        cell.border    = border
        cell.alignment = wrap

    # Data rows
    for idx, row in enumerate(rows, start=2):
        ws.append(row)
        row_type = row[1] if len(row) > 1 else None

        font_color = _C.NORMAL
        fill_color = None

        if row_colors and row_type in row_colors:
            font_color, fill_color = row_colors[row_type]
        elif row_type in (
            EntryType.EMPTY_BASE_OVERRIDE,
            EntryType.EMPTY_BASE_OVERRIDE_XML,
            EntryType.JSON_EMPTY_BASE_OVERRIDE,
            EntryType.LOGROTATE_EMPTY_BASE_OVERRIDE,
            EntryType.DUPLICATE_KEY,
            EntryType.INVALID_JSON,
            EntryType.INVALID_OUTPUT_JSON,
        ):
            font_color, fill_color = _C.RED, _C.RED_BG
        elif row_type == EntryType.BASE_ONLY_PARAMETER_ADDED:
            font_color, fill_color = _C.GREEN, _C.GREEN_BG
        elif row_type == EntryType.RELEASE_ONLY_PARAMETER_ADDED:
            font_color, fill_color = _C.ORANGE, _C.ORANGE_BG

        for cell in ws[idx]:
            cell.font      = Font(color=font_color)
            cell.border    = border
            cell.alignment = wrap
            if fill_color:
                cell.fill = PatternFill("solid", fgColor=fill_color)

    # Auto-size columns (capped at 80)
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                for line in str(cell.value).split("\n"):
                    max_len = max(max_len, len(line))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 80)


# ---------------------------------------------------------------------------
# Public writer
# ---------------------------------------------------------------------------

def write_excel(
    result: MergeResult,
    config: MergeConfig,
    report_dir: str = "reports",
    base_name: str = "",
) -> Path:
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    if not base_name:
        base_name = (
            result.base_name or
            os.path.basename(os.path.normpath(config.base_dir or "base"))
        )
    rel_name  = (
        os.path.basename(os.path.normpath(config.release_dirs[0]))
        if config.release_dirs else ""
    )

    # ── Sheet 1: MergeChanges ────────────────────────────────────────────
    changes_rows = []
    for entry in result.report:
        changes_rows.append([
            entry.file,
            entry.type,
            entry.element,
            entry.old,
            entry.new,
            _build_changes_column(entry),
        ])
    _add_sheet(wb, "MergeChanges",
               ["Merged File", "Merge Category", "Parameter",
                "Release Config Value", "Base/Prod Config Value", "Changes"],
               changes_rows)

    # ── Sheet 2: BaseOnlyFiles ───────────────────────────────────────────
    base_only_rows = [
        [f"{base_name}/{f}" if base_name else f]
        for f in sorted(result.base_only_files)
    ]
    _add_sheet(wb, "BaseOnlyFiles",
               ["Production Files (not in release)"],
               base_only_rows)

    # ── Sheet 3: ReleaseOnlyFiles ────────────────────────────────────────
    rel_only_rows = [
        [f"{rel_name}/{f}" if rel_name else f]
        for f in sorted(result.release_only_files)
    ]
    _add_sheet(wb, "ReleaseOnlyFiles",
               ["Release Config Files (not in base)"],
               rel_only_rows)

    # ── Sheet 4: FileMappings ────────────────────────────────────────────
    mapping_rows = [
        [f"{base_name}/{b}" if base_name else b,
         f"{rel_name}/{r}"  if rel_name  else r]
        for b, r in result.file_mappings
    ]
    # Flag multi-base mappings
    for rel_path, base_list in result.multi_base_mappings:
        for b in base_list:
            mapping_rows.append([
                f"{base_name}/{b} [MULTI-BASE]",
                f"{rel_name}/{rel_path}",
            ])
    _add_sheet(wb, "FileMappings",
               ["Base/Prod Config Filename", "Release Config Filename"],
               mapping_rows)

    # ── Sheet 5: ExcludedBaseOnly ────────────────────────────────────────
    excl_rows = [
        [e.file, e.element, e.new, e.base_comment]
        for e in result.excluded_params
    ]
    _add_sheet(wb, "ExcludedBaseOnly",
               ["File", "Parameter", "Base Value", "Base Comment"],
               excl_rows)

    # ── Sheet 6: BaseConfigAsIs ──────────────────────────────────────────
    copy_rows = [[f] for f in sorted(result.copy_only_files)]
    _add_sheet(wb, "BaseConfigAsIs",
               ["File copied as-is from Base"],
               copy_rows)

    Path(report_dir).mkdir(parents=True, exist_ok=True)
    # Filename encodes the base name so multiple bases land in the same dir
    # without colliding.  The directory itself carries the run timestamp.
    fname = f"merge_report_{base_name}.xlsx" if base_name else "merge_report.xlsx"
    report_path = Path(report_dir) / fname
    wb.save(report_path)
    return report_path
